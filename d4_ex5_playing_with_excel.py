# Task:
# Un pipeline de încarcare a datelor streaming dintr-un excel
# în care se interpune un validator bazat pe pydantic
# și care încarcă lucrurile în final în pandas.

# but first let's generate some fake data
import csv
from faker import Faker
import random

def mk_fake_data(csvout, rows=1000):
    faker = Faker('ro')

    with open(csvout, 'w') as f:
        writer = csv.DictWriter(f, ['Name','Email','Country'])
        writer.writeheader()

        for _ in range(rows):
            email = faker.email()
            code = faker.country_code()

            # but let's randomly mangle some data
            hmm = random.choices(["ok", "email", "code"], k=2, weights=[9, .4, .6])
            if 'email' in hmm:
                email = email.replace("@", "#")
            if 'code' in hmm:
                code = code + '_x'

            writer.writerow({
                'Name': faker.name(),
                'Email': email,
                'Country': code,
            })

#mk_fake_data('data/demo-dataset.csv')

##################
##   and now... ##
##################

## first, let's process the Excel file as a generator

import openpyxl

def xlsx_to_dicts(path):
    """
    Load data from an xlsx file and stream it as
    rows of dictionaries.
    """
    # Note: must open workbook with read_only=True
    #       in order to not load the spreadsheet in memory
    workbook = openpyxl.load_workbook(path,
                                      read_only=True,
                                      data_only=True)

    sheet = workbook[workbook.sheetnames[0]]
    rows = sheet.iter_rows()

    header = [cell.value for cell in next(rows)]

    for row in rows:
        yield dict(
            zip(
                header,
                (cell.value for cell in row)
            )
        )


## second, create a pydantic model
## and create a generator for it to consume the above

# must pip install pydantic[email]
#                  pydantic_extra_types
#              and pycountry

from pydantic import BaseModel, EmailStr, ValidationError
from pydantic_extra_types.country import CountryAlpha2

class Customer(BaseModel):
    Name: str
    Email: EmailStr
    Country: CountryAlpha2


def customer_processor(dataset, error_handler=None):
    """
    Validates dataset and yields Customer objects.

    To silence validation errors pass it an error_handler.
    """
    for row in dataset:
        try:
            # rows come in as dicts
            customer = Customer(**row)
        except ValidationError as e:
            if error_handler is not None:
                error_handler(e)
            else:
                raise e
        else:
            yield customer



## third, we can feed this directly do pandas:

import pandas

data1 = xlsx_to_dicts('data/demo-dataset.xlsx')
data2 = customer_processor(data1)
# everything is streaming so far
# so we get an early exit on errors:
#df = pandas.DataFrame.from_records(data2)

## or let's capture errors and create a dataframe anyway



errors = []
error_handler = errors.append
#error_handler = print
data2 = customer_processor(data1, error_handler)
df = pandas.DataFrame.from_records(data2)

print("### DataFrame is...")
print(df)

if errors:
    print("\n### But we have errors:")
    for idx, err in enumerate(errors):
        print(f'#{idx} ::', err)

        if idx == 10:
            break

    if len(errors) > 10:
        print("\n### %d more errors silenced" % (len(errors) - 10))